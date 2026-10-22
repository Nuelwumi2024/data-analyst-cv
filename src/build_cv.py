"""Build the CV as an ATS-friendly DOCX and a PDF from cv_content.py.

Single column, standard section headings, real text (no images or text boxes), A4 page.
Usage: python src/build_cv.py
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from reportlab.lib.colors import HexColor
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import HRFlowable, KeepTogether, ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer

import cv_content as c

OUT = Path(__file__).resolve().parents[1]
BASENAME = "Emmanuel_Omowumi_Data_Analyst_CV"
INK, ACCENT, MUTED = "1F2933", "1F4E79", "5F6B7A"


# ------------------------------------------------------------------ DOCX
def _border_bottom(paragraph, color=ACCENT, size=6):
    pPr = paragraph._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", str(size)), ("w:space", "1"), ("w:color", color)):
        bottom.set(qn(k), v)
    b.append(bottom)
    pPr.append(b)


def _hyperlink(paragraph, text, url, size=9.5):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color"); color.set(qn("w:val"), ACCENT)
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(int(size * 2)))
    rFonts = OxmlElement("w:rFonts"); rFonts.set(qn("w:ascii"), "Calibri"); rFonts.set(qn("w:hAnsi"), "Calibri")
    for el in (rFonts, color, sz):
        rPr.append(el)
    run.append(rPr)
    t = OxmlElement("w:t"); t.text = text; t.set(qn("xml:space"), "preserve")
    run.append(t); link.append(run)
    paragraph._p.append(link)


def build_docx(path: Path) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.8)
    sec.top_margin, sec.bottom_margin = Cm(1.4), Cm(1.3)
    base = doc.styles["Normal"]
    base.font.name, base.font.size = "Calibri", Pt(10)
    base.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    base.paragraph_format.space_after = Pt(0)
    base.paragraph_format.line_spacing = 1.05

    def para(text="", bold=False, size=10, color=INK, align=None, before=0, after=0, italic=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(before), Pt(after)
        if align:
            p.alignment = align
        if text:
            r = p.add_run(text); r.bold, r.italic, r.font.size = bold, italic, Pt(size)
            r.font.color.rgb = RGBColor.from_string(color)
        return p

    def heading(text):
        p = para(text.upper(), bold=True, size=10.5, color=ACCENT, before=7, after=3)
        _border_bottom(p)
        p.paragraph_format.keep_with_next = True

    def bullets(items):
        for it in items:
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(1.5)
            p.paragraph_format.left_indent = Cm(0.55)
            p.paragraph_format.first_line_indent = Cm(-0.4)
            r = p.add_run(it); r.font.size = Pt(9.5)

    para(c.NAME, bold=True, size=22, color=ACCENT, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(c.TITLE, size=10.5, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    p = para(align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    for i, (text, url) in enumerate(c.CONTACT):
        if i:
            r = p.add_run("  |  "); r.font.size = Pt(9.5); r.font.color.rgb = RGBColor.from_string(MUTED)
        if url:
            _hyperlink(p, text, url)
        else:
            r = p.add_run(text); r.font.size = Pt(9.5)

    heading("Professional summary")
    para(c.SUMMARY, size=9.5)

    heading("Core skills")
    for label, text in c.SKILLS:
        p = para(after=1)
        r = p.add_run(label + ": "); r.bold, r.font.size = True, Pt(9.5)
        r = p.add_run(text); r.font.size = Pt(9.5)

    heading("Professional experience")
    for job in c.EXPERIENCE:
        p = para(before=3)
        r = p.add_run(job["role"]); r.bold, r.font.size = True, Pt(10.5)
        r = p.add_run(f"  |  {job['org']}"); r.bold, r.font.size = True, Pt(10)
        p.paragraph_format.keep_with_next = True
        p = para(f"{job['dates']}  |  {job['where']}", size=9, color=MUTED, after=1, italic=True)
        p.paragraph_format.keep_with_next = True
        bullets(job["bullets"])

    heading("Selected projects")
    para(c.PORTFOLIO_LINE, size=9, color=MUTED, after=1, italic=True)
    for pr in c.PROJECTS:
        p = para(before=3)
        r = p.add_run(pr["name"]); r.bold, r.font.size = True, Pt(10)
        r = p.add_run(f"  |  {pr['stack']}"); r.font.size, r.italic = Pt(9), True
        r.font.color.rgb = RGBColor.from_string(MUTED)
        p.paragraph_format.keep_with_next = True
        bullets(pr["bullets"])
        p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.15)
        _hyperlink(p, pr["url"], "https://" + pr["url"], size=8.5)

    heading("Education")
    for deg, school, yrs in c.EDUCATION:
        p = para(after=1)
        r = p.add_run(deg); r.bold, r.font.size = True, Pt(9.5)
        r = p.add_run(f"  |  {school}  |  {yrs}"); r.font.size = Pt(9.5)

    heading("Training and certifications")
    bullets(c.CERTIFICATIONS)
    doc.core_properties.author, doc.core_properties.title = c.NAME, f"{c.NAME} - CV"
    doc.save(path)


# ------------------------------------------------------------------ PDF
def build_pdf(path: Path) -> None:
    ink, acc, mut = HexColor("#" + INK), HexColor("#" + ACCENT), HexColor("#" + MUTED)
    S = lambda name, **kw: ParagraphStyle(name, fontName=kw.pop("font", "Helvetica"), textColor=kw.pop("color", ink), **kw)
    st = {
        "name": S("name", font="Helvetica-Bold", fontSize=22, leading=26, alignment=TA_CENTER, color=acc),
        "title": S("title", fontSize=10.5, leading=13, alignment=TA_CENTER, color=mut),
        "contact": S("contact", fontSize=9, leading=12, alignment=TA_CENTER, spaceAfter=2),
        "h": S("h", font="Helvetica-Bold", fontSize=10.5, leading=13, color=acc, spaceBefore=8, spaceAfter=1),
        "body": S("body", fontSize=9.2, leading=12),
        "role": S("role", font="Helvetica-Bold", fontSize=10, leading=13, spaceBefore=4),
        "meta": S("meta", font="Helvetica-Oblique", fontSize=8.6, leading=11, color=mut, spaceAfter=1),
        "bul": S("bul", fontSize=9.2, leading=11.8),
        "url": S("url", fontSize=8.3, leading=10.5, color=acc),
    }
    esc = lambda t: t.replace("&", "&amp;").replace("<", "&lt;")
    story = [Paragraph(esc(c.NAME), st["name"]), Paragraph(esc(c.TITLE), st["title"])]
    parts = [f'<a href="{u}" color="#{ACCENT}">{esc(t)}</a>' if u else esc(t) for t, u in c.CONTACT]
    story.append(Paragraph("  |  ".join(parts), st["contact"]))

    def head(t):
        story.append(Paragraph(t.upper(), st["h"]))
        story.append(HRFlowable(width="100%", thickness=0.8, color=acc, spaceAfter=3))

    def bl(items):
        return ListFlowable([ListItem(Paragraph(esc(i), st["bul"]), leftIndent=11, bulletColor=ink) for i in items],
                            bulletType="bullet", start="•", leftIndent=11, bulletFontSize=8)

    head("Professional summary"); story.append(Paragraph(esc(c.SUMMARY), st["body"]))
    head("Core skills")
    for label, text in c.SKILLS:
        story.append(Paragraph(f"<b>{esc(label)}:</b> {esc(text)}", st["body"]))
    head("Professional experience")
    for j in c.EXPERIENCE:
        story.append(KeepTogether([
            Paragraph(f"{esc(j['role'])}  |  {esc(j['org'])}", st["role"]),
            Paragraph(f"{j['dates']}  |  {esc(j['where'])}", st["meta"]), bl(j["bullets"])]))
    head("Selected projects")
    story.append(Paragraph(esc(c.PORTFOLIO_LINE), st["meta"]))
    for pr in c.PROJECTS:
        story.append(KeepTogether([
            Paragraph(f"<b>{esc(pr['name'])}</b>  |  <i>{esc(pr['stack'])}</i>", st["role"]), bl(pr["bullets"]),
            Paragraph(f'<a href="https://{pr["url"]}" color="#{ACCENT}">{pr["url"]}</a>', st["url"])]))
    head("Education")
    for deg, school, yrs in c.EDUCATION:
        story.append(Paragraph(f"<b>{esc(deg)}</b>  |  {esc(school)}  |  {yrs}", st["body"]))
    head("Training and certifications"); story.append(bl(c.CERTIFICATIONS))
    SimpleDocTemplate(str(path), pagesize=A4, leftMargin=1.8 * cm, rightMargin=1.8 * cm, topMargin=1.3 * cm,
                      bottomMargin=1.2 * cm, title=f"{c.NAME} - CV", author=c.NAME).build(story)


if __name__ == "__main__":
    build_docx(OUT / f"{BASENAME}.docx")
    build_pdf(OUT / f"{BASENAME}.pdf")
    print("built", BASENAME)
